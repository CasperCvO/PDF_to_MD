"""Vision-based presentation-to-Markdown extraction (add-on module).

Adds multimodal extraction of presentations (PDF/PPTX) and pre-rendered slide
images (PNG/JPEG) into structured Markdown via the Google Gemini API
(``gemini-3.7-flash`` by default). This is an *add-on*: it does not replace or
modify the existing local converters (``convert_pdf``, ``convert_office``).

The module targets presentations that mix native text with embedded UI
screenshots, terminal dumps, and metric dashboards:

- Every slide is rendered to a high-resolution image (300 DPI by default) and
  sent to the vision model together with an extraction prompt.
- The model returns JSON (topic, detected technologies, Markdown body) which is
  written to a per-slide Markdown file with YAML frontmatter.
- An optional unified JSON export aggregates all slides for programmatic use.
- Requests run concurrently (configurable) with exponential-backoff retries
  for HTTP 429 / transient network errors.

Requirements
------------
- ``google-genai`` for the Gemini API.
- ``pillow`` for image validation and size normalization.
- ``pymupdf`` (already a dependency) for PDF page rendering.
- LibreOffice CLI (``soffice``) for ``.pptx`` rendering; if it is not
  installed, native slide text is extracted via ``python-pptx`` as a fallback
  (no vision analysis, ``vision_analyzed: false`` in the frontmatter).

API key resolution (in order): ``--api-key`` argument, the ``GEMINI_API_KEY``
environment variable, or a ``.env`` file in the current working directory.

Usage
-----
    python -m pdf_to_md.vision_converter --input-path slides.pdf --output-dir out

    python -m pdf_to_md.vision_converter --input-path deck.pptx --export-format both

    python -m pdf_to_md.vision_converter --input-path slide.png --api-key $GEMINI_API_KEY
"""

from __future__ import annotations

import argparse
import asyncio
import io
import json
import logging
import math
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

LOGGER = logging.getLogger(__name__)

API_KEY_ENV_VAR = "GEMINI_API_KEY"
DEFAULT_MODEL = "gemini-3.7-flash"
DEFAULT_DPI = 300
DEFAULT_CONCURRENCY = 4
DEFAULT_MAX_RETRIES = 5
DEFAULT_OUTPUT_DIR_NAME = "vision_extracts"

# Gemini inline-image limits: keep payloads safely below the documented
# ~20 MB / ~15 M pixel caps so we never trip a hard rejection.
MAX_IMAGE_BYTES = 19 * 1024 * 1024
MAX_IMAGE_PIXELS = 14_000_000
MIN_IMAGE_DIMENSION = 100
NATIVE_TEXT_HINT_CHARS = 6000

SUPPORTED_DOCUMENT_EXTENSIONS: set[str] = {".pdf", ".pptx"}
SUPPORTED_IMAGE_EXTENSIONS: set[str] = {".png", ".jpg", ".jpeg"}
MIME_BY_IMAGE_EXTENSION: dict[str, str] = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}

EXTRACTION_PROMPT = """\
You are an expert technical documentation analyst. You are given one slide image
from a technical presentation. Extract its complete content into
machine-readable GitHub-flavored Markdown for an LLM context repository. The
slides often contain native text, embedded UI screenshots, terminal dumps, and
metric dashboards - handle each faithfully.

Return ONLY a single raw JSON object (no Markdown fences) with exactly these keys:
- "topic": a short, precise title for the slide (the slide's own title if present).
- "technologies": an array of technology tags detected on the slide (e.g. "Pega",
  "PostgreSQL", "Kafka", "Alfresco", "PGPool", "Kubernetes"). Empty array if none.
- "markdown_body": the full slide content as Markdown.

Rules for "markdown_body":
1. Transcribe ALL visible text: titles, headings, body bullets, labels, axis
   labels, and notes.
2. Reconstruct tables (pod status tables, API response matrices, uptime lists,
   UI tables) as GitHub-flavored Markdown tables with a header row and separator.
3. Analyze charts and metric graphs (trendlines, SLA gauges, CPU/memory
   utilization spikes, bar charts) into concise numerical summaries as bullet
   points - state the metric, the values, and the trend (e.g. "CPU utilization
   peaked at 85% around 14:00, then returned to ~20%").
4. Extract terminal logs, shell commands, and monospace dumps verbatim into
   fenced code blocks tagged bash or text. Keep prompt symbols intact
   (e.g. "[user@host ~]$").
5. Use Markdown headings, bullet lists, bold labels, tables, and code blocks to
   preserve structure and reading order.
6. Never invent data that is not visible in the image. If text is illegible,
   mark it with "[illegible]".
"""


class _InvalidModelOutput(ValueError):
    """Raised when the model returns malformed JSON; safe to retry."""


@dataclass(frozen=True)
class SlidePage:
    """A single slide to process.

    Either *image_bytes* (vision path) or *native_text* (fallback path) is
    populated; the other may be None.
    """

    slide_number: int
    image_bytes: bytes | None = None
    mime_type: str | None = None
    native_text: str | None = None


@dataclass
class SlideResult:
    """Extraction result for a single slide."""

    slide_number: int
    topic: str
    technologies: list[str]
    markdown_body: str
    vision_analyzed: bool


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def convert_presentation_vision(
    input_path: Path | str,
    output_dir: Path | str | None = None,
    *,
    api_key: str | None = None,
    model: str = DEFAULT_MODEL,
    dpi: int = DEFAULT_DPI,
    export_format: str = "md",
    concurrency: int = DEFAULT_CONCURRENCY,
    max_retries: int = DEFAULT_MAX_RETRIES,
    page_start: int | None = None,
    page_end: int | None = None,
    extraction_prompt: str | None = None,
) -> list[Path]:
    """Extract Markdown from a presentation or slide images using Gemini vision.

    Parameters
    ----------
    input_path : Path | str
        A ``.pdf``/``.pptx`` presentation, a ``.png``/``.jpg``/``.jpeg`` slide
        image, or a directory containing any of those (processed in sorted
        order).
    output_dir : Path | str | None, optional
        Directory for the generated files. If None, defaults to
        ``<input_path.parent>/vision_extracts`` (or ``<input_path>/vision_extracts``
        when *input_path* is a directory).
    api_key : str | None, optional
        Gemini API key. If None, the ``GEMINI_API_KEY`` environment variable is
        used. Only required when slides are analyzed with vision (never for the
        LibreOffice-less pptx text fallback).
    model : str, optional
        Gemini model id (default ``gemini-3.7-flash``).
    dpi : int, optional
        Render resolution for PDF/PPTX pages (default 300).
    export_format : str, optional
        One of ``"md"`` (per-slide Markdown files, default), ``"json"`` (one
        unified JSON array), or ``"both"``.
    concurrency : int, optional
        Maximum number of simultaneous API requests (default 4).
    max_retries : int, optional
        Retries per slide for HTTP 429/5xx, network errors, and malformed model
        output (default 5, exponential backoff).
    page_start, page_end : int | None, optional
        1-based inclusive slide range to process.
    extraction_prompt : str | None, optional
        Override the built-in extraction prompt.

    Returns
    -------
    list[Path]
        Paths of all files written (Markdown and/or JSON).
    """
    _validate_options(export_format, dpi, concurrency, max_retries, page_start, page_end, model)

    input_path = Path(input_path).resolve()
    if not input_path.exists():
        raise FileNotFoundError(f"Input path not found: {input_path}")

    if output_dir is None:
        output_dir = (
            input_path / DEFAULT_OUTPUT_DIR_NAME
            if input_path.is_dir()
            else input_path.parent / DEFAULT_OUTPUT_DIR_NAME
        )
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    effective_prompt = extraction_prompt.strip() if extraction_prompt else EXTRACTION_PROMPT
    api_key = _resolve_api_key(api_key)

    if input_path.is_dir():
        files = sorted(
            (
                path
                for path in input_path.iterdir()
                if path.suffix.lower()
                in (SUPPORTED_DOCUMENT_EXTENSIONS | SUPPORTED_IMAGE_EXTENSIONS)
            ),
            key=lambda p: p.name.lower(),
        )
        if not files:
            raise FileNotFoundError(
                f"No supported files (.pdf, .pptx, .png, .jpg, .jpeg) found in {input_path}"
            )
    else:
        files = [input_path]

    written: list[Path] = []
    for file_path in files:
        try:
            written.extend(
                _convert_one(
                    file_path,
                    output_dir,
                    api_key=api_key,
                    model=model,
                    dpi=dpi,
                    export_format=export_format,
                    concurrency=concurrency,
                    max_retries=max_retries,
                    page_start=page_start,
                    page_end=page_end,
                    extraction_prompt=effective_prompt,
                )
            )
        except Exception as exc:  # noqa: BLE001 - batch mode keeps going
            if input_path.is_dir():
                LOGGER.error("Failed to process %s: %s", file_path, exc)
            else:
                raise
    return written


# ---------------------------------------------------------------------------
# Loading / rendering
# ---------------------------------------------------------------------------


def _load_document_slides(
    file_path: Path,
    dpi: int,
    page_start: int | None = None,
    page_end: int | None = None,
) -> list[SlidePage]:
    """Load slides from a PDF or PPTX; PPTX prefers LibreOffice rendering.

    When *page_start*/*page_end* are given, only that 1-based inclusive range is
    loaded so huge documents are not fully rendered for a partial run.
    """
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return _render_pdf_pages(file_path, dpi, page_start, page_end)
    if ext == ".pptx":
        try:
            with _pptx_to_pdf(file_path) as pdf_path:
                return _render_pdf_pages(pdf_path, dpi, page_start, page_end)
        except FileNotFoundError as exc:
            LOGGER.warning(
                "%s; falling back to native text extraction for %s (no vision analysis).",
                exc,
                file_path.name,
            )
            return _pptx_text_slides(file_path, page_start, page_end)
    raise ValueError(f"Unsupported document extension: {ext}")


def _render_pdf_pages(
    pdf_path: Path,
    dpi: int,
    page_start: int | None = None,
    page_end: int | None = None,
) -> list[SlidePage]:
    """Render PDF pages to high-resolution PNGs (plus native-text hints).

    Only pages within the 1-based inclusive [*page_start*, *page_end*] range are
    rendered; page numbers are preserved as document page numbers.
    """
    try:
        import pymupdf
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency. Install pymupdf to render PDF pages."
        ) from exc

    pages: list[SlidePage] = []
    with pymupdf.open(pdf_path) as doc:
        if doc.page_count == 0:
            raise ValueError(f"PDF has no pages: {pdf_path}")
        first = max(1, page_start or 1)
        last = min(doc.page_count, page_end or doc.page_count)
        if first > last:
            return pages
        matrix = pymupdf.Matrix(dpi / 72.0, dpi / 72.0)
        for number in range(first, last + 1):
            page = doc.load_page(number - 1)
            native_text = page.get_text("text").strip() or None
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            data = pix.tobytes("png")
            data, mime_type = _normalize_image_bytes(data, "image/png", pix.width, pix.height)
            pages.append(
                SlidePage(
                    slide_number=number,
                    image_bytes=data,
                    mime_type=mime_type,
                    native_text=native_text,
                )
            )
    return pages


def _load_single_image(image_path: Path) -> SlidePage:
    """Validate a raw slide image and return it as a one-slide payload."""
    mime_type, data = _validate_and_prepare_image(image_path)
    return SlidePage(slide_number=1, image_bytes=data, mime_type=mime_type)


def _pptx_text_slides(
    pptx_path: Path,
    page_start: int | None = None,
    page_end: int | None = None,
) -> list[SlidePage]:
    """Extract native slide text from a PPTX (fallback when LibreOffice is absent).

    Only slides within the 1-based inclusive [*page_start*, *page_end*] range
    are extracted; slide numbers are preserved.
    """
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency. Install python-pptx for the .pptx text fallback."
        ) from exc

    prs = Presentation(str(pptx_path))
    if not prs.slides:
        raise ValueError(f"PowerPoint file has no slides: {pptx_path}")

    pages: list[SlidePage] = []
    for number, slide in enumerate(prs.slides, start=1):
        if page_start is not None and number < page_start:
            continue
        if page_end is not None and number > page_end:
            continue
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        blocks.append(text)
            elif getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if rows:
                    header = "| " + " | ".join(rows[0]) + " |"
                    separator = "|" + "|".join("---" for _ in rows[0]) + "|"
                    body = ["| " + " | ".join(row) + " |" for row in rows[1:]]
                    blocks.append("\n".join([header, separator, *body]))
        native_text = "\n\n".join(blocks).strip() or "(no extractable text found)"
        pages.append(
            SlidePage(slide_number=number, image_bytes=None, mime_type=None, native_text=native_text)
        )
    return pages


def _find_soffice() -> str | None:
    """Locate the LibreOffice CLI on PATH or at common Windows install paths."""
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    if sys.platform == "win32":
        candidates = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            os.path.expandvars(r"%LOCALAPPDATA%\Programs\LibreOffice\program\soffice.exe"),
        ]
        for candidate in candidates:
            if Path(candidate).is_file():
                return candidate
    return None


@contextmanager
def _pptx_to_pdf(pptx_path: Path) -> Iterator[Path]:
    """Headlessly convert a PPTX to PDF via LibreOffice (temp dir auto-cleaned)."""
    soffice = _find_soffice()
    if soffice is None:
        raise FileNotFoundError(
            "LibreOffice (soffice) is required to render .pptx files to images, but it was "
            "not found. Install LibreOffice (https://www.libreoffice.org/) or pre-convert "
            "the file to PDF."
        )
    with tempfile.TemporaryDirectory(prefix="pdf_to_md_vision_") as tmp_dir:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp_dir, str(pptx_path)],
            capture_output=True,
            text=True,
            timeout=600,
        )
        if result.returncode != 0:
            raise RuntimeError(
                f"LibreOffice conversion failed: {(result.stderr or result.stdout).strip()}"
            )
        pdf_path = Path(tmp_dir) / f"{pptx_path.stem}.pdf"
        if not pdf_path.is_file():
            raise RuntimeError("LibreOffice reported success but produced no PDF.")
        yield pdf_path


# ---------------------------------------------------------------------------
# Image validation / normalization
# ---------------------------------------------------------------------------


def _validate_and_prepare_image(image_path: Path) -> tuple[str, bytes]:
    """Validate *image_path* (mime, resolution, color depth) and return (mime, bytes).

    Raises ``ValueError`` for files that are not usable PNG/JPEG images.
    """
    try:
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency. Install pillow to enable image validation."
        ) from exc

    expected_mime = MIME_BY_IMAGE_EXTENSION.get(image_path.suffix.lower())
    if expected_mime is None:
        raise ValueError(
            f"Unsupported image extension: {image_path.suffix} (supported: .png, .jpg, .jpeg)"
        )

    try:
        with Image.open(image_path) as img:
            actual_format = (img.format or "").upper()
            if actual_format not in ("PNG", "JPEG"):
                raise ValueError(
                    f"Image validation failed for {image_path.name}: declared {expected_mime}, "
                    f"detected {actual_format or 'unknown'} format."
                )
            width, height = img.size
            if width < MIN_IMAGE_DIMENSION or height < MIN_IMAGE_DIMENSION:
                raise ValueError(
                    f"Image too small for vision analysis: {width}x{height} "
                    f"(minimum {MIN_IMAGE_DIMENSION}px per side)."
                )
            if img.mode not in ("RGB", "RGBA", "L", "LA"):
                LOGGER.info(
                    "Converting %s from color mode %s to RGB.", image_path.name, img.mode
                )
                img = img.convert("RGB")
            if img.mode not in ("RGB", "RGBA", "L", "LA"):
                raise ValueError(f"Unsupported color depth for {image_path.name}: {img.mode}")
            save_format = "PNG" if expected_mime == "image/png" else "JPEG"
            buffer = io.BytesIO()
            img.save(buffer, format=save_format)
            data = buffer.getvalue()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Image validation failed for {image_path.name}: {exc}") from exc

    data, mime_type = _normalize_image_bytes(data, expected_mime, width, height)
    return mime_type, data


def _normalize_image_bytes(data: bytes, mime_type: str, width: int, height: int) -> tuple[bytes, str]:
    """Downscale/re-encode *data* so it fits Gemini's inline image limits.

    Returns the payload unchanged when it is already within limits.
    """
    if len(data) <= MAX_IMAGE_BYTES and width * height <= MAX_IMAGE_PIXELS:
        return data, mime_type
    try:
        from PIL import Image
    except ImportError:
        LOGGER.warning(
            "pillow not available; sending %d-byte image unchanged.", len(data)
        )
        return data, mime_type

    with Image.open(io.BytesIO(data)) as img:
        if width * height > MAX_IMAGE_PIXELS:
            scale = math.sqrt(MAX_IMAGE_PIXELS / (width * height))
            img = img.resize(
                (max(1, int(width * scale)), max(1, int(height * scale))),
                Image.LANCZOS,
            )
        img = img.convert("RGB")
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=88)
        normalized = buffer.getvalue()
    LOGGER.info(
        "Downscaled/re-encoded slide image to fit API limits (%d -> %d bytes).",
        len(data),
        len(normalized),
    )
    return normalized, "image/jpeg"


# ---------------------------------------------------------------------------
# Vision extraction
# ---------------------------------------------------------------------------


def _resolve_api_key(api_key: str | None) -> str | None:
    """Resolve the API key from (in order): argument, env var, ``.env`` file.

    A ``.env`` file in the current working directory is loaded as a fallback
    (KEY=VALUE lines; existing environment variables are never overwritten).
    """
    if api_key:
        return api_key
    api_key = os.environ.get(API_KEY_ENV_VAR)
    if api_key:
        return api_key
    _load_dotenv()
    return os.environ.get(API_KEY_ENV_VAR)


def _load_dotenv(path: Path | None = None) -> None:
    """Load KEY=VALUE pairs from a ``.env`` file without overwriting the environment.

    Only simple lines are supported: ``KEY=VALUE``, optional surrounding quotes,
    and ``#`` comments. The key is never logged.
    """
    env_path = path if path is not None else Path.cwd() / ".env"
    if not env_path.is_file():
        return
    for line in env_path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


def _create_async_client(api_key: str):
    """Create an async Gemini client (lazy import keeps the add-on optional)."""
    from google import genai

    return genai.Client(api_key=api_key).aio


async def _extract_all_async(
    pages: list[SlidePage],
    api_key: str,
    model: str,
    document_source: str,
    extraction_prompt: str,
    concurrency: int,
    max_retries: int,
) -> list[SlideResult]:
    """Run vision extraction for *pages* with bounded concurrency and retries."""
    client = _create_async_client(api_key)
    try:
        return await _process_slides(
            client, model, pages, document_source, extraction_prompt, concurrency, max_retries
        )
    finally:
        await client.aclose()


async def _process_slides(
    client,
    model: str,
    pages: list[SlidePage],
    document_source: str,
    extraction_prompt: str,
    concurrency: int,
    max_retries: int,
) -> list[SlideResult]:
    """Fan out per-slide extraction across a semaphore-bounded pool."""
    semaphore = asyncio.Semaphore(concurrency)

    async def worker(page: SlidePage) -> SlideResult:
        async with semaphore:
            return await _extract_slide_with_retry(
                client, model, page, document_source, extraction_prompt, max_retries
            )

    return await asyncio.gather(*(worker(page) for page in pages))


async def _extract_slide_with_retry(
    client,
    model: str,
    page: SlidePage,
    document_source: str,
    extraction_prompt: str,
    max_retries: int,
) -> SlideResult:
    """Extract one slide, retrying transient failures with exponential backoff."""
    last_error: BaseException | None = None
    attempts_made = 0
    for attempt in range(max_retries + 1):
        attempts_made = attempt + 1
        try:
            result = await _extract_slide_once(client, model, page, document_source, extraction_prompt)
            LOGGER.info(
                "Extracted slide %d (%s) with %d technology tag(s).",
                page.slide_number,
                result.topic or "untitled",
                len(result.technologies),
            )
            return result
        except Exception as exc:  # noqa: BLE001 - retry classification below
            last_error = exc
            if not _is_retryable(exc) or attempt == max_retries:
                break
            delay = min(60.0, 2**attempt) + random.uniform(0.0, 0.5)
            LOGGER.warning(
                "Slide %d attempt %d/%d failed (%s). Retrying in %.1fs.",
                page.slide_number,
                attempt + 1,
                max_retries + 1,
                _describe_error(exc),
                delay,
            )
            await asyncio.sleep(delay)
    raise RuntimeError(
        f"Slide {page.slide_number} failed after {attempts_made} attempt(s) "
        f"(max_retries={max_retries}): {_describe_error(last_error)}"
    ) from last_error


async def _extract_slide_once(
    client,
    model: str,
    page: SlidePage,
    document_source: str,
    extraction_prompt: str,
) -> SlideResult:
    """Send one slide image to the vision model and parse the JSON response."""
    prompt = _build_slide_prompt(extraction_prompt, page, document_source)
    contents = [
        {
            "role": "user",
            "parts": [
                {"text": prompt},
                {"inline_data": {"mime_type": page.mime_type, "data": page.image_bytes}},
            ],
        }
    ]
    response = await asyncio.wait_for(
        client.models.generate_content(
            model=model,
            contents=contents,
            config={"response_mime_type": "application/json"},
        ),
        timeout=180,
    )

    text = getattr(response, "text", None)
    if not text:
        feedback = getattr(response, "prompt_feedback", None)
        reason = getattr(feedback, "block_reason", None) if feedback else None
        raise _InvalidModelOutput(f"Model returned no text (block reason: {reason})")

    payload = _parse_model_json(text)
    topic = str(payload.get("topic", "") or "").strip()
    technologies = [
        str(tag).strip() for tag in payload.get("technologies", []) if str(tag).strip()
    ]
    markdown_body = str(payload.get("markdown_body", "") or "").strip()
    return SlideResult(
        slide_number=page.slide_number,
        topic=topic,
        technologies=technologies,
        markdown_body=markdown_body,
        vision_analyzed=True,
    )


def _build_slide_prompt(base_prompt: str, page: SlidePage, document_source: str) -> str:
    """Append per-slide context (document, number, native-text hint) to the prompt."""
    context = f"\n\nDocument: {document_source}\nSlide number: {page.slide_number}"
    if page.native_text:
        hint = page.native_text[:NATIVE_TEXT_HINT_CHARS]
        context += (
            "\n\nNative text extracted from the document's text layer is provided below as a "
            "transcription aid. Verify it against the rendered image and correct any layout/"
            "OCR mistakes; do not duplicate it verbatim.\n\n"
            f"--- native text ---\n{hint}\n--- end native text ---"
        )
    return base_prompt + context


def _parse_model_json(text: str) -> dict:
    """Parse the model's JSON response, tolerating stray Markdown fences."""
    cleaned = text.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```[a-zA-Z]*\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start == -1 or end <= start:
            raise _InvalidModelOutput(f"Model did not return valid JSON: {text[:200]!r}") from None
        try:
            data = json.loads(cleaned[start : end + 1])
        except json.JSONDecodeError as exc:
            raise _InvalidModelOutput(
                f"Model did not return valid JSON: {text[:200]!r}"
            ) from exc
    if not isinstance(data, dict):
        raise _InvalidModelOutput(f"Model returned a non-object JSON value: {type(data).__name__}")
    return data


def _describe_error(exc: BaseException | None) -> str:
    """Human-readable one-liner for API errors (falls back to str())."""
    if exc is None:
        return "unknown error"
    message = getattr(exc, "message", None)
    status = getattr(exc, "status", None)
    code = getattr(exc, "code", None)
    if message and status and code:
        return f"{code} {status}: {message}"
    return str(exc)


def _is_retryable(exc: BaseException) -> bool:
    """Return True for HTTP 429/5xx, network errors, and malformed model output."""
    if isinstance(exc, _InvalidModelOutput):
        return True
    code = getattr(exc, "code", None)
    if isinstance(code, str):
        try:
            code = int(code)
        except ValueError:
            code = None
    if isinstance(code, int) and code in (429, 500, 502, 503, 504):
        return True
    if isinstance(exc, (TimeoutError, ConnectionError)):
        return True
    try:
        import httpx
    except ImportError:
        return False
    return isinstance(exc, httpx.TransportError)


# ---------------------------------------------------------------------------
# Output writers
# ---------------------------------------------------------------------------


def _write_slide_markdown(
    result: SlideResult,
    document_source: str,
    output_dir: Path,
    clean_stem: str,
) -> Path:
    """Write one self-contained Markdown file with YAML frontmatter."""
    # json.dumps yields valid double-quoted YAML scalars for every value.
    frontmatter = "\n".join(
        [
            "---",
            f"document_source: {json.dumps(document_source, ensure_ascii=False)}",
            f"slide_number: {result.slide_number}",
            f"topic: {json.dumps(result.topic, ensure_ascii=False)}",
            f"technologies: {json.dumps(result.technologies, ensure_ascii=False)}",
            f"vision_analyzed: {str(result.vision_analyzed).lower()}",
            "---",
            "",
        ]
    )
    body = result.markdown_body.strip() or "*No content extracted.*"
    content = f"{frontmatter}\n{body}\n"
    path = output_dir / f"{clean_stem}_slide_{result.slide_number:03d}.md"
    path.write_text(content, encoding="utf-8")
    return path


def _result_to_dict(result: SlideResult, document_source: str) -> dict:
    """Serialize a slide result for the unified JSON export."""
    return {
        "document_source": document_source,
        "slide_number": result.slide_number,
        "topic": result.topic,
        "technologies": result.technologies,
        "markdown": result.markdown_body,
        "vision_analyzed": result.vision_analyzed,
    }


def _fallback_result(page: SlidePage) -> SlideResult:
    """Build a text-only result for slides that could not be vision-analyzed."""
    body = page.native_text or ""
    topic = ""
    if body:
        first_line = body.splitlines()[0].strip()
        if len(first_line) <= 120:
            topic = first_line
    return SlideResult(
        slide_number=page.slide_number,
        topic=topic,
        technologies=[],
        markdown_body=body,
        vision_analyzed=False,
    )


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


def _convert_one(
    file_path: Path,
    output_dir: Path,
    *,
    api_key: str | None,
    model: str,
    dpi: int,
    export_format: str,
    concurrency: int,
    max_retries: int,
    page_start: int | None,
    page_end: int | None,
    extraction_prompt: str,
) -> list[Path]:
    """Process a single input file and return the paths of written outputs."""
    ext = file_path.suffix.lower()
    if ext in SUPPORTED_DOCUMENT_EXTENSIONS:
        slides = _load_document_slides(file_path, dpi, page_start, page_end)
    elif ext in SUPPORTED_IMAGE_EXTENSIONS:
        slides = [_load_single_image(file_path)]
    else:
        raise ValueError(
            f"Unsupported input format: {ext} "
            f"(supported: {sorted(SUPPORTED_DOCUMENT_EXTENSIONS | SUPPORTED_IMAGE_EXTENSIONS)})"
        )

    slides = _apply_page_range(slides, page_start, page_end)
    if not slides:
        raise ValueError(f"No slides remain after applying the page range for {file_path}")

    vision_pages = [slide for slide in slides if slide.image_bytes is not None]
    document_source = file_path.name

    if vision_pages:
        if not api_key:
            raise RuntimeError(
                f"No API key available. Set the {API_KEY_ENV_VAR} environment variable "
                "or pass --api-key."
            )
        LOGGER.info(
            "Extracting %d slide(s) from %s with model %s (concurrency=%d)...",
            len(vision_pages),
            document_source,
            model,
            concurrency,
        )
        vision_results = asyncio.run(
            _extract_all_async(
                vision_pages,
                api_key,
                model,
                document_source,
                extraction_prompt,
                concurrency,
                max_retries,
            )
        )
    else:
        vision_results = []

    results_by_number = {result.slide_number: result for result in vision_results}
    results = [
        results_by_number.get(slide.slide_number) or _fallback_result(slide)
        for slide in slides
    ]

    clean_stem = re.sub(r"[\s\(\)\[\]]+", "_", file_path.stem)
    written: list[Path] = []

    if export_format in ("md", "both"):
        for result in results:
            path = _write_slide_markdown(result, document_source, output_dir, clean_stem)
            written.append(path)
            LOGGER.info("Wrote %s (slide %d)", path.name, result.slide_number)

    if export_format in ("json", "both"):
        json_path = output_dir / f"{clean_stem}.json"
        entries = [_result_to_dict(result, document_source) for result in results]
        json_path.write_text(
            json.dumps(entries, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        written.append(json_path)
        LOGGER.info("Wrote %s", json_path.name)

    return written


def _apply_page_range(
    slides: list[SlidePage], page_start: int | None, page_end: int | None
) -> list[SlidePage]:
    """Filter slides to the 1-based inclusive [page_start, page_end] range."""
    if page_start is None and page_end is None:
        return slides
    start = page_start if page_start is not None else 1
    end = page_end if page_end is not None else max(slide.slide_number for slide in slides)
    return [slide for slide in slides if start <= slide.slide_number <= end]


def _validate_options(
    export_format: str,
    dpi: int,
    concurrency: int,
    max_retries: int,
    page_start: int | None,
    page_end: int | None,
    model: str,
) -> None:
    """Validate user-facing options and raise ValueError with clear messages."""
    if export_format not in ("md", "json", "both"):
        raise ValueError(f"export_format must be 'md', 'json', or 'both' (got {export_format!r})")
    if not 72 <= dpi <= 600:
        raise ValueError(f"dpi must be between 72 and 600 (got {dpi})")
    if concurrency < 1:
        raise ValueError(f"concurrency must be >= 1 (got {concurrency})")
    if max_retries < 0:
        raise ValueError(f"max_retries must be >= 0 (got {max_retries})")
    if page_start is not None and page_start < 1:
        raise ValueError(f"page_start must be >= 1 (got {page_start})")
    if page_end is not None and page_end < 1:
        raise ValueError(f"page_end must be >= 1 (got {page_end})")
    if page_start is not None and page_end is not None and page_start > page_end:
        raise ValueError(f"page_start ({page_start}) must not exceed page_end ({page_end})")
    if not model.strip():
        raise ValueError("model must not be empty")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Extract structured Markdown from presentations (PDF/PPTX) and slide images "
            "(PNG/JPEG) using Gemini vision."
        ),
    )
    # Note: the --input-path option uses its own dest because an optional flag
    # sharing dest with a nargs="?" positional is silently discarded by argparse.
    parser.add_argument(
        "input_path",
        nargs="?",
        default=None,
        help="Presentation, slide image, or directory to process.",
    )
    parser.add_argument(
        "--input-path",
        dest="input_path_opt",
        default=None,
        help="Same as the positional input_path.",
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        default=None,
        help=(
            f"Output directory (default: <input_path>/vision_extracts or "
            f"<input_path parent>/vision_extracts)."
        ),
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=DEFAULT_DPI,
        help=f"Render resolution for PDF/PPTX pages (default: {DEFAULT_DPI}).",
    )
    parser.add_argument(
        "--export-format",
        choices=["md", "json", "both"],
        default="md",
        help=(
            "md: per-slide Markdown files; json: one unified JSON array; "
            "both: write both (default: md)."
        ),
    )
    parser.add_argument(
        "--api-key",
        default=None,
        help=f"Gemini API key (default: the {API_KEY_ENV_VAR} environment variable).",
    )
    parser.add_argument(
        "--model",
        default=DEFAULT_MODEL,
        help=f"Gemini model id (default: {DEFAULT_MODEL}).",
    )
    parser.add_argument(
        "--concurrency",
        type=int,
        default=DEFAULT_CONCURRENCY,
        help=f"Max simultaneous API requests (default: {DEFAULT_CONCURRENCY}).",
    )
    parser.add_argument(
        "--max-retries",
        type=int,
        default=DEFAULT_MAX_RETRIES,
        help=(
            "Retries per slide for HTTP 429/5xx, network errors, and malformed "
            f"output (default: {DEFAULT_MAX_RETRIES})."
        ),
    )
    parser.add_argument(
        "--page-start",
        type=int,
        default=None,
        help="First slide to process (1-based, inclusive).",
    )
    parser.add_argument(
        "--page-end",
        type=int,
        default=None,
        help="Last slide to process (1-based, inclusive).",
    )
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    # Keep third-party SDK noise (httpx request lines, AFC notices/warnings) out of the CLI.
    for noisy_logger in ("httpx", "httpcore"):
        logging.getLogger(noisy_logger).setLevel(logging.WARNING)
    logging.getLogger("google_genai").setLevel(logging.ERROR)

    input_path = args.input_path or args.input_path_opt
    if not input_path:
        parser.error("input_path is required (positional or --input-path)")

    input_resolved = Path(input_path).resolve()
    try:
        written = convert_presentation_vision(
            input_path,
            args.output_dir,
            api_key=args.api_key,
            model=args.model,
            dpi=args.dpi,
            export_format=args.export_format,
            concurrency=args.concurrency,
            max_retries=args.max_retries,
            page_start=args.page_start,
            page_end=args.page_end,
        )
    except Exception as exc:
        LOGGER.error("Conversion failed: %s", exc)
        raise SystemExit(1) from exc
    output_dir = (
        Path(args.output_dir).resolve()
        if args.output_dir
        else (
            input_resolved / DEFAULT_OUTPUT_DIR_NAME
            if input_resolved.is_dir()
            else input_resolved.parent / DEFAULT_OUTPUT_DIR_NAME
        )
    )
    LOGGER.info("Wrote %d file(s) to %s", len(written), output_dir)


if __name__ == "__main__":
    main()
