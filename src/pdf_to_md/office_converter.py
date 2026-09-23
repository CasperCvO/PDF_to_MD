"""Office-document-to-Markdown converter module.

Converts office files (docx, xlsx, pptx, xls, doc, …) to Markdown using
Microsoft's *markitdown* library — except ``.xlsx`` workbooks, which are
converted with openpyxl so every sheet gets a ``<!-- sheet: "Name"
range: A1:D4 -->`` locator and a row/column-addressed table for provenance.

``.docx`` files additionally get approximate ``<!-- page: N -->`` markers
(on by default, disable with ``docx_page_markers=False``): token runs are
injected into ``word/document.xml`` at Word's saved layout breaks
(``w:lastRenderedPageBreak``) or, when absent, at manual page breaks,
page-break-before paragraphs and page-starting section breaks, then
rewritten into markers after conversion.

Each output file is prepended with a Wikilink to the original document for
easy back-referencing in Obsidian.
"""

from __future__ import annotations

import copy
import datetime
import io
import logging
import re
import zipfile
from pathlib import Path
from typing import Any

LOGGER = logging.getLogger(__name__)

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W_BODY = f"{{{_W_NS}}}body"
_W_P = f"{{{_W_NS}}}p"
_W_PPR = f"{{{_W_NS}}}pPr"
_W_R = f"{{{_W_NS}}}r"
_W_RPR = f"{{{_W_NS}}}rPr"
_W_T = f"{{{_W_NS}}}t"
_W_BR = f"{{{_W_NS}}}br"
_W_LRPB = f"{{{_W_NS}}}lastRenderedPageBreak"
_W_PBB = f"{{{_W_NS}}}pageBreakBefore"
_W_SECTPR = f"{{{_W_NS}}}sectPr"
_W_TYPE = f"{{{_W_NS}}}type"
_W_VAL = f"{{{_W_NS}}}val"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
_SKIPPED_ANCESTORS = {f"{{{_W_NS}}}txbxContent", f"{{{_W_NS}}}del"}

_PAGE_TOKEN = r"PDFTOMDPAGE\d+END"
_PAGE_TOKEN_NUMBER = re.compile(r"PDFTOMDPAGE(\d+)END")
_PAGE_TOKEN_LINE = re.compile(
    rf"^(#{{1,6}} )?((?:{_PAGE_TOKEN}[ \t]*)+)([^\n]*)", re.MULTILINE
)
_PAGE_TOKEN_INLINE = re.compile(rf"[ \t]*(?:{_PAGE_TOKEN}[ \t]*)+")

SUPPORTED_EXTENSIONS: set[str] = {
    ".docx",
    ".doc",
    ".xlsx",
    ".xls",
    ".pptx",
    ".ppt",
}


def _format_cell(value: Any) -> str:
    """Render a single cell value as Markdown-safe text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str):
        text = value.strip()
        text = (
            text.replace("\r\n", "<br>").replace("\r", "<br>").replace("\n", "<br>")
        )
        return text.replace("|", "\\|")
    return str(value)


def _xlsx_to_markdown(file_path: Path) -> str:
    """Return Markdown for an ``.xlsx`` workbook with sheet/cell locators."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency. Install openpyxl to enable .xlsx conversion."
        ) from exc

    formulas_wb = load_workbook(file_path, data_only=False)
    values_wb = load_workbook(file_path, data_only=True)

    sheets = []
    for ws in formulas_wb.worksheets:
        values_ws = values_wb[ws.title]
        non_empty: set[tuple[int, int]] = set()
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                value = cell.value
                if value is not None and (
                    not isinstance(value, str) or value.strip()
                ):
                    non_empty.add((cell.row, cell.column))

        heading = f"## Sheet: {ws.title}"
        if not non_empty:
            sheets.append(f"{heading}\n\n_(empty sheet)_")
            continue

        rows = sorted({row for row, _ in non_empty})
        cols = sorted({col for _, col in non_empty})
        min_row, max_row = rows[0], rows[-1]
        min_col, max_col = cols[0], cols[-1]
        top_left = f"{get_column_letter(min_col)}{min_row}"
        bottom_right = f"{get_column_letter(max_col)}{max_row}"
        escaped_title = ws.title.replace('"', '\\"')
        locator = f'<!-- sheet: "{escaped_title}" range: {top_left}:{bottom_right} -->'

        header = "| Row | " + " | ".join(get_column_letter(col) for col in cols) + " |"
        separator = "| --- |" + " --- |" * len(cols)
        lines = [header, separator]
        for row in rows:
            cells = []
            for col in cols:
                if (row, col) not in non_empty:
                    cells.append("")
                    continue
                value = values_ws.cell(row=row, column=col).value
                if value is None:
                    formula = ws.cell(row=row, column=col).value
                    value = getattr(formula, "text", None) or str(formula)
                cells.append(_format_cell(value))
            lines.append(f"| {row} | " + " | ".join(cells) + " |")

        sheets.append(f"{heading}\n\n{locator}\n\n" + "\n".join(lines))

    return "\n\n".join(sheets)


def _has_skipped_ancestor(element) -> bool:
    """Return True when *element* sits inside a ``w:txbxContent`` or ``w:del``."""
    return any(
        ancestor.tag in _SKIPPED_ANCESTORS
        for ancestor in element.iterancestors()
    )


def _docx_page_break_candidates(body) -> list[tuple[str, Any]]:
    """Collect ``(kind, element)`` page-break candidates in document order.

    Rendered mode: when any usable ``w:lastRenderedPageBreak`` exists, only
    those are returned — Word already emitted one after every explicit break.
    Explicit mode: manual ``w:br`` page breaks, ``w:pageBreakBefore``
    paragraphs and page-starting ``w:sectPr`` section ends (whose break lands
    on the next usable paragraph).
    """
    rendered = [
        element
        for element in body.iter(_W_LRPB)
        if not _has_skipped_ancestor(element)
    ]
    if rendered:
        return [("break", element) for element in rendered]

    candidates: list[tuple[str, Any]] = []
    pending_section_break = False
    first_paragraph_seen = False
    for element in body.iter():
        if not isinstance(element.tag, str) or _has_skipped_ancestor(element):
            continue
        if element.tag == _W_P:
            is_first = not first_paragraph_seen
            first_paragraph_seen = True
            if pending_section_break:
                candidates.append(("start", element))
                pending_section_break = False
            else:
                properties = element.find(_W_PPR)
                page_break_before = (
                    properties.find(_W_PBB) if properties is not None else None
                )
                if (
                    page_break_before is not None
                    and page_break_before.get(_W_VAL) not in {"0", "false", "off"}
                    and not is_first
                ):
                    candidates.append(("start", element))
            properties = element.find(_W_PPR)
            section = properties.find(_W_SECTPR) if properties is not None else None
            if section is not None:
                section_type = section.find(_W_TYPE)
                value = (
                    section_type.get(_W_VAL) if section_type is not None else None
                )
                if value is None or value in {"nextPage", "oddPage", "evenPage"}:
                    pending_section_break = True
        elif element.tag == _W_BR and element.get(_W_TYPE) == "page":
            candidates.append(("break", element))
    return candidates


def _make_page_token_run(token: str):
    """Build an unstyled ``w:r`` run whose text is *token*."""
    from lxml import etree

    run = etree.Element(_W_R)
    text = etree.SubElement(run, _W_T)
    text.set(_XML_SPACE, "preserve")
    text.text = token
    return run


def _insert_token_at_break(break_element, token: str) -> bool:
    """Split the run containing *break_element* and put *token* between the parts."""
    from lxml import etree

    run = break_element.getparent()
    if run is None or run.tag != _W_R:
        return False

    children = list(run)
    index = children.index(break_element)
    run_properties = run.find(_W_RPR)

    second = etree.Element(_W_R)
    if run_properties is not None:
        second.append(copy.deepcopy(run_properties))
    for child in children[index + 1:]:
        second.append(child)
    run.remove(break_element)

    parent = run.getparent()
    position = list(parent).index(run)
    parent.insert(position + 1, _make_page_token_run(token))
    if any(child.tag != _W_RPR for child in second):
        parent.insert(position + 2, second)
    return True


def _insert_token_at_paragraph_start(paragraph, token: str) -> bool:
    """Insert *token* as the first run of *paragraph*."""
    properties = paragraph.find(_W_PPR)
    index = list(paragraph).index(properties) + 1 if properties is not None else 0
    paragraph.insert(index, _make_page_token_run(token))
    return True


def _insert_docx_page_tokens(docx_bytes: bytes) -> tuple[bytes, int]:
    """Inject ``PDFTOMDPAGE<n>END`` token runs at page-break positions.

    Returns the (possibly unchanged) document bytes and the number of tokens
    inserted.  Token *n* is the page that starts at that position, so the
    first break produces ``PDFTOMDPAGE2END``.
    """
    from lxml import etree

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as source:
        try:
            document_xml = source.read("word/document.xml")
        except KeyError:
            return docx_bytes, 0
        root = etree.fromstring(document_xml)
        body = root.find(_W_BODY)
        if body is None:
            return docx_bytes, 0

        candidates = _docx_page_break_candidates(body)
        if not candidates:
            return docx_bytes, 0

        inserted = 0
        for kind, element in candidates:
            token = f"PDFTOMDPAGE{inserted + 2}END"
            if kind == "break":
                success = _insert_token_at_break(element, token)
            else:
                success = _insert_token_at_paragraph_start(element, token)
            if success:
                inserted += 1

        if not inserted:
            return docx_bytes, 0

        new_xml = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "word/document.xml":
                    data = new_xml
                target.writestr(info, data)

    return buffer.getvalue(), inserted


def _replace_page_tokens(markdown: str) -> str:
    """Rewrite ``PDFTOMDPAGE<n>END`` tokens as ``<!-- page: n -->`` markers."""

    def line_sub(match: re.Match) -> str:
        heading = match.group(1) or ""
        numbers = _PAGE_TOKEN_NUMBER.findall(match.group(2))
        markers = "\n\n".join(f"<!-- page: {n} -->" for n in numbers)
        rest = match.group(3)
        if rest.strip():
            return f"{markers}\n\n{heading}{rest.lstrip()}"
        return markers

    def inline_sub(match: re.Match) -> str:
        numbers = _PAGE_TOKEN_NUMBER.findall(match.group(0))
        markers = " ".join(f"<!-- page: {n} -->" for n in numbers)
        text = match.string
        lead = "" if match.start() == 0 or text[match.start() - 1] == "\n" else " "
        trail = "" if match.end() == len(text) or text[match.end()] == "\n" else " "
        return f"{lead}{markers}{trail}"

    return _PAGE_TOKEN_INLINE.sub(
        inline_sub, _PAGE_TOKEN_LINE.sub(line_sub, markdown)
    )


def _docx_to_markdown(file_path: Path) -> str:
    """Return Markdown for a ``.docx`` with ``<!-- page: N -->`` markers."""
    try:
        from markitdown import MarkItDown, StreamInfo
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency. Install markitdown to enable office-document conversion: "
            "pip install 'markitdown[docx,xlsx,pptx]'"
        ) from exc

    docx_bytes, inserted = _insert_docx_page_tokens(file_path.read_bytes())
    converter = MarkItDown()
    result = converter.convert_stream(
        io.BytesIO(docx_bytes), stream_info=StreamInfo(extension=".docx")
    )
    markdown = _replace_page_tokens((result.text_content or "").strip())
    body = f"<!-- page: 1 -->\n\n{markdown}" if markdown else "<!-- page: 1 -->"

    found = body.count("<!-- page:")
    if found != inserted + 1:
        LOGGER.warning(
            "Page markers lost converting %s: expected %d, found %d",
            file_path,
            inserted + 1,
            found,
        )
    return body


def _extract_markdown(file_path: Path, *, docx_page_markers: bool = True) -> str:
    """Return Markdown text for *file_path*."""
    suffix = file_path.suffix.lower()
    if suffix == ".xlsx":
        return _xlsx_to_markdown(file_path)
    if suffix == ".docx" and docx_page_markers:
        return _docx_to_markdown(file_path)

    try:
        from markitdown import MarkItDown
    except ImportError as exc:
        raise RuntimeError(
            "Missing dependency. Install markitdown to enable office-document conversion: "
            "pip install 'markitdown[docx,xlsx,pptx]'"
        ) from exc

    converter = MarkItDown()
    result = converter.convert(str(file_path))
    return (result.text_content or "").strip()


def convert_office(
    file_path: Path,
    output_path: Path,
    *,
    docx_page_markers: bool = True,
) -> None:
    """Convert a single office document to Markdown and write to *output_path*.

    When *docx_page_markers* is true, ``.docx`` output gets approximate
    ``<!-- page: N -->`` markers (Word's saved layout breaks when present,
    else manual/page-break-before/section breaks).
    """
    markdown_body = _extract_markdown(file_path, docx_page_markers=docx_page_markers)
    wikilink = f"[[{file_path.name}]]"
    content = f"{wikilink}\n\n{markdown_body}\n" if markdown_body else f"{wikilink}\n"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(content, encoding="utf-8")
    LOGGER.info("Wrote %s", output_path)
