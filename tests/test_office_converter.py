"""Tests for XLSX sheet/cell locators and DOCX page markers in office_converter."""

import sys
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from pdf_to_md.office_converter import _replace_page_tokens, convert_office

EXPECTED_DATA_BLOCK = '''## Sheet: Data

<!-- sheet: "Data" range: A1:D4 -->

| Row | A | B | D |
| --- | --- | --- | --- |
| 1 | Name | Amount |  |
| 2 | Alice | 100 |  |
| 4 | Bob\\|Smith | =SUM(B2:B2) | x |'''


def _write_test_workbook(path: Path) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Amount"
    ws["A2"] = "Alice"
    ws["B2"] = 100
    ws["A4"] = "Bob|Smith"
    ws["B4"] = "=SUM(B2:B2)"
    ws["D4"] = "x"

    wb.create_sheet("Notes")

    offset = wb.create_sheet("Offset")
    offset["C5"] = 2.5
    offset["D6"] = datetime(2024, 1, 31)

    misc = wb.create_sheet("Misc")
    misc["A1"] = "line one\nline two"
    misc["B1"] = True

    wb.save(path)


def test_xlsx_markdown(tmp_path):
    xlsx_path = tmp_path / "book.xlsx"
    _write_test_workbook(xlsx_path)
    output_path = tmp_path / "out" / "book.md"

    convert_office(xlsx_path, output_path)
    text = output_path.read_text(encoding="utf-8")

    assert text.splitlines()[0] == "[[book.xlsx]]"
    assert EXPECTED_DATA_BLOCK in text
    assert '<!-- sheet: "Offset" range: C5:D6 -->' in text
    assert "| 5 | 2.5 |  |" in text
    assert "| 6 |  | 2024-01-31 |" in text
    assert "## Sheet: Notes\n\n_(empty sheet)_" in text
    assert "| 3 |" not in text
    assert "line one<br>line two" in text
    assert "| TRUE |" in text


_DOCX_CONTENT_TYPES = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
</Types>'''

_DOCX_RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''

_DOCX_DOCUMENT_RELS = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>'''

_DOCX_STYLES = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/></w:style>
</w:styles>'''


def _write_docx(path: Path, body_xml: str) -> None:
    """Zip a minimal .docx package whose ``w:body`` content is *body_xml*."""
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:body>'
        + body_xml
        + '</w:body></w:document>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _DOCX_CONTENT_TYPES)
        zf.writestr("_rels/.rels", _DOCX_RELS)
        zf.writestr("word/_rels/document.xml.rels", _DOCX_DOCUMENT_RELS)
        zf.writestr("word/styles.xml", _DOCX_STYLES)
        zf.writestr("word/document.xml", document)


RENDERED_BODY = """
<w:p><w:r><w:t>Intro on page one</w:t></w:r></w:p>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:lastRenderedPageBreak/><w:t>Chapter Two</w:t></w:r></w:p>
<w:p><w:r><w:t xml:space="preserve">Sentence starts on page two </w:t></w:r><w:r><w:lastRenderedPageBreak/><w:t>and ends on page three.</w:t></w:r></w:p>
<w:p><w:r><w:br w:type="page"/></w:r></w:p>
<w:p><w:r><w:rPr><w:b/></w:rPr><w:lastRenderedPageBreak/><w:t>Bold page four</w:t></w:r></w:p>
<w:tbl>
<w:tr><w:tc><w:p><w:r><w:t>Cell A</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Cell B</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:lastRenderedPageBreak/><w:t>Cell C</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Cell D</w:t></w:r></w:p></w:tc></w:tr>
</w:tbl>
<w:p><w:r><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:lastRenderedPageBreak/><w:t>box</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>
"""

EXPLICIT_BODY = """
<w:p><w:pPr><w:pageBreakBefore/></w:pPr><w:r><w:t>First page</w:t></w:r></w:p>
<w:p><w:r><w:t>Before break</w:t><w:br w:type="page"/><w:t>After break</w:t></w:r></w:p>
<w:p><w:pPr><w:pageBreakBefore/></w:pPr><w:r><w:t>Chapter on page three</w:t></w:r></w:p>
<w:p><w:pPr><w:pageBreakBefore w:val="0"/></w:pPr><w:r><w:t>Still page three</w:t></w:r></w:p>
<w:p><w:pPr><w:sectPr><w:type w:val="nextPage"/></w:sectPr></w:pPr><w:r><w:t>End of section</w:t></w:r></w:p>
<w:p><w:r><w:t>New section page four</w:t></w:r></w:p>
<w:p><w:pPr><w:sectPr><w:type w:val="continuous"/></w:sectPr></w:pPr><w:r><w:t>Continuous end</w:t></w:r></w:p>
<w:p><w:r><w:t>Same page four</w:t></w:r></w:p>
<w:sectPr/>
"""


def test_docx_rendered_page_markers(tmp_path):
    docx_path = tmp_path / "rendered.docx"
    _write_docx(docx_path, RENDERED_BODY)
    output_path = tmp_path / "out" / "rendered.md"

    convert_office(docx_path, output_path)
    text = output_path.read_text(encoding="utf-8")

    lines = text.splitlines()
    assert lines[0] == f"[[{docx_path.name}]]"
    assert lines[2] == "<!-- page: 1 -->"
    assert "<!-- page: 2 -->\n\n# Chapter Two" in text
    assert "page two <!-- page: 3 --> and ends on page three." in text
    assert "<!-- page: 4 -->\n\n**Bold page four**" in text
    assert "| <!-- page: 5 --> Cell C |" in text
    assert text.count("<!-- page:") == 5
    assert "PDFTOMDPAGE" not in text


def test_docx_explicit_page_markers(tmp_path):
    docx_path = tmp_path / "explicit.docx"
    _write_docx(docx_path, EXPLICIT_BODY)
    output_path = tmp_path / "out" / "explicit.md"

    convert_office(docx_path, output_path)
    text = output_path.read_text(encoding="utf-8")

    assert text.count("<!-- page:") == 4
    assert "Before break <!-- page: 2 --> After break" in text
    assert "<!-- page: 3 -->\n\nChapter on page three" in text
    assert "<!-- page: 4 -->\n\nNew section page four" in text
    assert "Still page three" in text
    assert "Same page four" in text
    assert "PDFTOMDPAGE" not in text


def test_docx_page_markers_disabled(tmp_path):
    from markitdown import MarkItDown

    docx_path = tmp_path / "rendered.docx"
    _write_docx(docx_path, RENDERED_BODY)
    output_path = tmp_path / "out" / "rendered.md"

    convert_office(docx_path, output_path, docx_page_markers=False)
    text = output_path.read_text(encoding="utf-8")

    expected = (
        f"[[{docx_path.name}]]\n\n"
        + (MarkItDown().convert(str(docx_path)).text_content or "").strip()
        + "\n"
    )
    assert text == expected
    assert "<!-- page:" not in text


def test_convert_tree_docx_page_markers(tmp_path):
    from pdf_to_md import convert_tree

    source = tmp_path / "src"
    source.mkdir()
    _write_docx(source / "doc.docx", RENDERED_BODY)

    out_marked = tmp_path / "marked"
    convert_tree(source, out_marked)
    assert "<!-- page:" in (out_marked / "doc.md").read_text(encoding="utf-8")

    out_plain = tmp_path / "plain"
    convert_tree(source, out_plain, docx_page_markers=False)
    assert "<!-- page:" not in (out_plain / "doc.md").read_text(encoding="utf-8")


def test_cli_docx_page_markers_flag(tmp_path, monkeypatch):
    import pdf_to_md.main as main_module

    recorded = {}

    def _recorder(source_root, output_root=None, *, docx_page_markers=True):
        recorded["docx_page_markers"] = docx_page_markers
        return {"pdf": 0, "office": 0, "skipped": 0}

    monkeypatch.setattr(main_module, "convert_tree", _recorder)
    source = tmp_path / "src"
    source.mkdir()

    monkeypatch.setattr(
        sys, "argv", ["pdf-to-md", str(source), "--no-docx-page-markers"]
    )
    main_module.main()
    assert recorded["docx_page_markers"] is False

    monkeypatch.setattr(sys, "argv", ["pdf-to-md", str(source)])
    main_module.main()
    assert recorded["docx_page_markers"] is True


def test_replace_page_tokens_unit():
    assert (
        _replace_page_tokens("PDFTOMDPAGE3END**Bold** text")
        == "<!-- page: 3 -->\n\n**Bold** text"
    )
    assert (
        _replace_page_tokens("# PDFTOMDPAGE3ENDTitle")
        == "<!-- page: 3 -->\n\n# Title"
    )
    assert (
        _replace_page_tokens("PDFTOMDPAGE2ENDPDFTOMDPAGE3END")
        == "<!-- page: 2 -->\n\n<!-- page: 3 -->"
    )
    assert (
        _replace_page_tokens("thePDFTOMDPAGE2END continued")
        == "the <!-- page: 2 --> continued"
    )
    assert (
        _replace_page_tokens("| PDFTOMDPAGE4ENDCell C | Cell D |")
        == "| <!-- page: 4 --> Cell C | Cell D |"
    )
    assert (
        _replace_page_tokens("* PDFTOMDPAGE5ENDItem")
        == "* <!-- page: 5 --> Item"
    )


def test_docx_empty_body(tmp_path):
    docx_path = tmp_path / "empty.docx"
    _write_docx(docx_path, "")
    output_path = tmp_path / "out" / "empty.md"

    convert_office(docx_path, output_path)
    text = output_path.read_text(encoding="utf-8")

    assert text == f"[[{docx_path.name}]]\n\n<!-- page: 1 -->\n"
